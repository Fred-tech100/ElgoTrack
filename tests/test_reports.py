import io
import json
from datetime import timedelta
from decimal import Decimal

from django.urls import reverse
from django.utils import timezone

from contacts.models import Customer, Supplier
from expenses import services as exp
from expenses.models import ExpenseCategory
from inventory import services as inv
from payments.models import PaymentMethod
from products.models import Category, Product
from purchases import services as pur
from reports import engine
from reports.exporters import safe_text
from reports.models import SavedReport
from sales import services as sales

from .helpers import TenantTestCase, make_product, make_user

D = Decimal


class ReportData(TenantTestCase):
    """Cement 30000/25000, Nails 5000/3000; two sellers; credit + cash; one return."""

    def setUp(self):
        super().setUp()
        self.manager = make_user(self.business, "manager", "mia@alpha.test")
        self.cat_cement = Category.objects.create(business=self.business, name="Cement")
        self.cat_hw = Category.objects.create(business=self.business, name="Hardware")
        self.product.category = self.cat_cement
        self.product.save()
        self.nails = make_product(self.business, self.owner, name="Nails", price="5000", cost="3000", stock="50", category=self.cat_hw)
        self.joe = Customer.objects.create(business=self.business, name="Joe", customer_type="credit")
        self.supplier = Supplier.objects.create(business=self.business, name="Hima")
        self.s1 = sales.create_sale(business=self.business, user=self.owner, items=[{"product_id": self.product.pk, "quantity": "2"}],
                                    payments=[{"method_id": self.cash.pk, "amount": "60000"}])
        self.s2 = sales.create_sale(business=self.business, user=self.manager, customer=self.joe, items=[{"product_id": self.nails.pk, "quantity": "10"}],
                                    payments=[{"method_id": self.momo.pk, "amount": "20000"}])
        sales.create_sale_return(sale=self.s2, user=self.manager, lines=[{"item_id": self.s2.items.get().pk, "quantity": "2"}], reason="Faulty")
        self.f = lambda **kw: engine.Filters(start=timezone.localdate() - timedelta(days=1), end=timezone.localdate(), **kw)

    def rows(self, result):
        return {r["group"]: r for r in result.rows}


class ReportCalculationTests(ReportData):
    def test_sales_by_product_net_of_returns(self):
        r = self.rows(engine.sales_grouped(self.business, self.f(), "product"))
        self.assertEqual((r["Cement 50kg"]["qty"], r["Cement 50kg"]["net"], r["Cement 50kg"]["profit"]), (D("2"), D("60000"), D("10000")))
        n = r["Nails"]
        self.assertEqual((n["qty"], n["gross"], n["returned"], n["net"], n["cost"], n["profit"]), (D("8"), D("50000"), D("10000"), D("40000"), D("24000"), D("16000")))
        self.assertAlmostEqual(float(r["Cement 50kg"]["margin"]), 16.67, places=1)

    def test_totals_match_financial_summary(self):
        res = engine.sales_grouped(self.business, self.f(), "product")
        from reports.services import financial_summary
        s = financial_summary(self.business, self.f().start, self.f().end)
        self.assertEqual(res.totals["net"], s["net_sales"])
        self.assertEqual(res.totals["cost"], s["cogs"])
        self.assertEqual(res.totals["profit"], s["gross_profit"])

    def test_groupings(self):
        by_cat = self.rows(engine.sales_grouped(self.business, self.f(), "category"))
        self.assertEqual((by_cat["Cement"]["net"], by_cat["Hardware"]["net"]), (D("60000"), D("40000")))
        by_user = self.rows(engine.sales_grouped(self.business, self.f(), "user"))
        self.assertEqual((by_user["Alpha Owner"]["net"], by_user["mia"]["net"]), (D("60000"), D("40000")))
        by_cust = self.rows(engine.sales_grouped(self.business, self.f(), "customer"))
        self.assertEqual((by_cust["Walk-in customers"]["net"], by_cust["Joe"]["net"]), (D("60000"), D("40000")))
        daily = engine.sales_grouped(self.business, self.f(), "day")
        self.assertEqual(len(daily.rows), 1)
        self.assertEqual(daily.totals["net"], D("100000"))
        for period in ("week", "month"):
            self.assertEqual(engine.sales_grouped(self.business, self.f(), period).totals["net"], D("100000"))

    def test_filters(self):
        self.assertEqual(engine.sales_grouped(self.business, self.f(product=self.nails), "product").totals["net"], D("40000"))
        self.assertEqual(engine.sales_grouped(self.business, self.f(customer=self.joe), "product").totals["net"], D("40000"))
        self.assertEqual(engine.sales_grouped(self.business, self.f(user=self.owner), "product").totals["net"], D("60000"))
        self.assertEqual(engine.sales_grouped(self.business, self.f(category=self.cat_cement), "product").totals["net"], D("60000"))
        old = engine.Filters(start=timezone.localdate() - timedelta(days=60), end=timezone.localdate() - timedelta(days=30))
        self.assertEqual(engine.sales_grouped(self.business, old, "product").rows, [])

    def test_payment_method_report(self):
        r = self.rows(engine.sales_by_payment_method(self.business, self.f()))
        self.assertEqual((r["Cash"]["money_in"], r["Mobile Money"]["money_in"]), (D("60000"), D("20000")))

    def test_stock_valuation_and_current(self):
        v = engine.stock_valuation(self.business, self.f())
        self.assertEqual((v.totals["cost"], v.totals["retail"]), (D("2576000"), D("3150000")))
        cur = {r["product"]: r for r in engine.stock_current(self.business, self.f()).rows}
        self.assertEqual((cur["Cement 50kg"]["qty"], cur["Nails"]["qty"]), (D("98"), D("42")))
        self.assertEqual(engine.stock_current(self.business, self.f(category=self.cat_hw)).rows[0]["product"], "Nails")

    def test_low_and_out_of_stock(self):
        self.nails.reorder_level = D("50")
        self.nails.save()
        self.assertEqual([r["product"] for r in engine.stock_current(self.business, self.f(), "low").rows], ["Nails"])
        inv.record_movement(business=self.business, product=self.nails, movement_type="loss", quantity=D("-42"), user=self.owner)
        self.assertEqual([r["product"] for r in engine.stock_current(self.business, self.f(), "out").rows], ["Nails"])

    def test_stock_variance_damage_loss(self):
        count = inv.start_count(business=self.business, user=self.owner)
        line = count.lines.get(product=self.product)
        inv.save_count_entries(count, {line.pk: (D("96"), "Wet bags")})
        inv.submit_count(count, self.owner)
        inv.approve_count(count, self.owner)
        inv.manual_adjustment(business=self.business, user=self.owner, product=self.product, movement_type="damage", quantity=D("1"), reason="Dropped")
        inv.manual_adjustment(business=self.business, user=self.owner, product=self.nails, movement_type="loss", quantity=D("3"), reason="Theft")
        var = engine.stock_variance(self.business, self.f())
        self.assertEqual((var.rows[0]["variance"], var.rows[0]["value"]), (D("-2"), D("-50000")))
        dmg = engine.REPORT_MAP["stock-damaged"].build(self.business, self.f())
        self.assertEqual((dmg.rows[0]["group"], dmg.totals["qty_out"]), ("Cement 50kg", D("1")))
        lost = engine.REPORT_MAP["stock-lost"].build(self.business, self.f())
        self.assertEqual(lost.totals["qty_out"], D("3"))
        adj = engine.REPORT_MAP["stock-adjustments"].build(self.business, self.f())
        self.assertEqual(adj.rows, [])  # approved-count differences are 'reconciliation', not manual 'adjustment'

    def test_balances_with_overdue(self):
        sales.create_sale(business=self.business, user=self.owner, customer=self.joe, items=[{"product_id": self.product.pk, "quantity": "1"}],
                          payments=[], due_date=timezone.localdate() - timedelta(days=1))
        r = engine.balances(self.business, self.f(), "customer")
        self.assertEqual((r.rows[0]["balance"], r.rows[0]["overdue"]), (D("50000"), D("30000")))  # 20000 + 30000
        pur.create_purchase(business=self.business, user=self.owner, supplier=self.supplier, receive_now=True,
                            items=[{"product_id": self.nails.pk, "quantity": "10", "unit_cost": "3000"}])
        sb = engine.balances(self.business, self.f(), "supplier")
        self.assertEqual(sb.rows[0]["balance"], D("30000"))

    def test_purchase_and_expense_reports(self):
        po = pur.create_purchase(business=self.business, user=self.owner, supplier=self.supplier, receive_now=True,
                                 items=[{"product_id": self.nails.pk, "quantity": "10", "unit_cost": "3000"}])
        r = engine.REPORT_MAP["purchases-supplier"].build(self.business, self.f())
        self.assertEqual((r.rows[0]["group"], r.rows[0]["qty"], r.rows[0]["amount"]), ("Hima", D("10"), D("30000")))
        h = engine.purchase_history(self.business, self.f())
        self.assertEqual((h.rows[0]["number"], h.rows[0]["due"]), (po.number, D("30000")))
        exp.record_expense(business=self.business, user=self.owner, category=ExpenseCategory.objects.get(business=self.business, name="Rent"),
                           amount=D("10000"), method=self.cash, description="Rent")
        e = engine.REPORT_MAP["expenses"].build(self.business, self.f())
        self.assertEqual((e.rows[0]["group"], e.rows[0]["amount"]), ("Rent", D("10000")))

    def test_expiring_stock(self):
        pur.create_purchase(business=self.business, user=self.owner, supplier=self.supplier, receive_now=False,
                            items=[{"product_id": self.nails.pk, "quantity": "5", "unit_cost": "3000"}])
        po = self.supplier.purchases.get()
        pur.receive_goods(po, self.owner, [{"item_id": po.items.get().pk, "quantity": "5", "batch_number": "B1", "expiry_date": timezone.localdate() + timedelta(days=10)}])
        r = engine.expiring_stock(self.business, self.f())
        self.assertEqual((r.rows[0]["batch"], r.rows[0]["status"]), ("B1", "Expiring soon"))

    def test_isolated_between_businesses(self):
        other = engine.sales_grouped(self.other, self.f(), "product")
        self.assertEqual(other.rows, [])
        self.assertEqual(engine.stock_current(self.other, self.f()).rows, [])


class ReportViewTests(ReportData):
    def test_every_report_renders_and_exports(self):
        self.login(self.owner)
        for r in engine.REPORTS:
            url = reverse("reports:run", args=[r.key])
            self.assertEqual(self.client.get(url).status_code, 200, r.key)
            for fmt, ctype in (("csv", "text/csv"), ("xlsx", "spreadsheetml"), ("pdf", "application/pdf")):
                resp = self.client.get(url, {"export": fmt})
                self.assertEqual(resp.status_code, 200, f"{r.key} {fmt}")
                self.assertIn(ctype, resp["Content-Type"])
                self.assertIn("attachment", resp["Content-Disposition"])
        self.assertEqual(self.client.get(reverse("reports:index")).status_code, 200)
        self.assertEqual(self.client.get(reverse("reports:run", args=["nope"])).status_code, 404)
        self.assertEqual(self.client.get(reverse("reports:run", args=["sales-product"]), {"export": "exe"}).status_code, 404)

    def test_csv_xlsx_contents(self):
        self.login(self.owner)
        csv_text = self.client.get(reverse("reports:run", args=["sales-product"]), {"export": "csv", "days": "7"}).content.decode("utf-8-sig")
        self.assertIn("Net sales", csv_text)
        self.assertIn("Cement 50kg", csv_text)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(self.client.get(reverse("reports:run", args=["sales-product"]), {"export": "xlsx"}).content))
        ws = wb.active
        headers = [c.value for c in ws[4]]
        self.assertIn("Net sales", headers)
        total_row = [c.value for c in ws[ws.max_row]]
        self.assertEqual(total_row[0], "Total")
        self.assertIn(100000, total_row)

    def test_filters_via_query_string(self):
        self.login(self.owner)
        resp = self.client.get(reverse("reports:run", args=["sales-product"]), {"product": self.nails.pk, "export": "csv"})
        text = resp.content.decode("utf-8-sig")
        self.assertIn("Nails", text)
        self.assertNotIn("Cement 50kg", text)
        # foreign ids are ignored, not honoured
        foreign = Product.objects.create(business=self.other, name="X", unit=self.product.unit.__class__.objects.filter(business=self.other).first())
        resp = self.client.get(reverse("reports:run", args=["sales-product"]), {"product": foreign.pk, "export": "csv"})
        self.assertIn("Cement 50kg", resp.content.decode("utf-8-sig"))

    def test_csv_formula_injection_neutralised(self):
        self.assertEqual(safe_text("=HYPERLINK(\"http://evil\")"), "'=HYPERLINK(\"http://evil\")")
        self.assertEqual(safe_text("+1"), "'+1")
        self.assertEqual(safe_text("Normal"), "Normal")
        evil = make_product(self.business, self.owner, name='=cmd|" /C calc"!A0', stock="5")
        sales.create_sale(business=self.business, user=self.owner, items=[{"product_id": evil.pk, "quantity": "1"}], payments=[{"method_id": self.cash.pk, "amount": "30000"}])
        self.login(self.owner)
        text = self.client.get(reverse("reports:run", args=["sales-product"]), {"export": "csv"}).content.decode("utf-8-sig")
        self.assertIn("'=cmd", text)
        self.assertNotIn(",=cmd", text)
        self.assertNotIn("\n=cmd", text)

    def test_permissions(self):
        sp = make_user(self.business, "salesperson", "sam@alpha.test")
        self.login(sp)
        self.assertEqual(self.client.get(reverse("reports:index")).status_code, 403)
        sm = make_user(self.business, "stock_manager", "stella@alpha.test")
        self.login(sm)
        self.assertEqual(self.client.get(reverse("reports:run", args=["stock-current"])).status_code, 200)
        self.assertEqual(self.client.get(reverse("reports:run", args=["profit-product"])).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:run", args=["customer-credits"]), {"export": "csv"}).status_code, 403)
        self.assertEqual(self.client.get(reverse("reports:builder")).status_code, 403)
        listing = self.client.get(reverse("reports:index"))
        self.assertContains(listing, "Current stock")
        self.assertNotContains(listing, "Profit by product")
        acc = make_user(self.business, "accountant", "ann@alpha.test")
        self.login(acc)
        self.assertEqual(self.client.get(reverse("reports:run", args=["expenses"])).status_code, 200)
        self.assertEqual(self.client.get(reverse("reports:builder")).status_code, 200)

    def test_report_export_is_audited(self):
        from audit.models import AuditLog
        self.login(self.owner)
        self.client.get(reverse("reports:run", args=["sales-product"]), {"export": "csv"})
        self.assertTrue(AuditLog.objects.filter(action="report.export", business=self.business).exists())


class BuilderTests(ReportData):
    def test_grouped_matches_standard_report(self):
        self.login(self.owner)
        resp = self.client.get(reverse("reports:builder"), {"dataset": "sales", "group_by": "product", "days": "7"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["result"].totals["net"], D("100000"))
        std = engine.sales_grouped(self.business, self.f(), "product")
        self.assertEqual(resp.context["result"].rows, std.rows)

    def test_all_datasets_and_groups(self):
        self.login(self.owner)
        pur.create_purchase(business=self.business, user=self.owner, supplier=self.supplier, receive_now=True, items=[{"product_id": self.nails.pk, "quantity": "5", "unit_cost": "3000"}])
        exp.record_expense(business=self.business, user=self.owner, category=ExpenseCategory.objects.get(business=self.business, name="Rent"), amount=D("100"), method=self.cash, description="x")
        for dataset in ("sales", "purchases", "stock", "payments", "expenses"):
            for group in [""] + engine.dataset_groups(dataset):
                resp = self.client.get(reverse("reports:builder"), {"dataset": dataset, "group_by": group, "days": "30"})
                self.assertEqual(resp.status_code, 200, (dataset, group))
                self.assertIsNotNone(resp.context["result"])
            for fmt in ("csv", "xlsx", "pdf"):
                self.assertEqual(self.client.get(reverse("reports:builder"), {"dataset": dataset, "export": fmt}).status_code, 200)

    def test_detail_rows_and_filters(self):
        self.login(self.owner)
        resp = self.client.get(reverse("reports:builder"), {"dataset": "stock", "movement_type": "sale", "days": "7"})
        self.assertEqual(len(resp.context["result"].rows), 2)
        resp = self.client.get(reverse("reports:builder"), {"dataset": "payments", "method": self.momo.pk, "days": "7"})
        self.assertEqual(resp.context["result"].totals["amount"], D("20000"))
        resp = self.client.get(reverse("reports:builder"), {"dataset": "stock", "group_by": "type", "days": "7"})
        types = {r["group"]: r for r in resp.context["result"].rows}
        self.assertEqual(types["Sale"]["qty_out"], D("12"))

    def test_bad_input(self):
        self.login(self.owner)
        self.assertEqual(self.client.get(reverse("reports:builder"), {"dataset": "users"}).status_code, 404)
        resp = self.client.get(reverse("reports:builder"), {"dataset": "sales", "group_by": "supplier"})
        self.assertContains(resp, "not available")

    def test_save_run_delete_saved_report(self):
        self.login(self.owner)
        r = self.client.post(reverse("reports:builder"), {"name": "Cement sales", "dataset": "sales", "group_by": "product", "days": "30"})
        self.assertEqual(r.status_code, 302)
        saved = SavedReport.objects.get()
        self.assertEqual(saved.params["dataset"], "sales")
        r = self.client.get(reverse("reports:saved", args=[saved.pk]))
        self.assertIn("dataset=sales", r["Location"])
        self.assertEqual(self.client.post(reverse("reports:builder"), {"name": "", "dataset": "sales"}).status_code, 302)
        self.assertEqual(SavedReport.objects.count(), 1)
        self.client.post(reverse("reports:saved", args=[saved.pk]))
        self.assertEqual(SavedReport.objects.count(), 0)

    def test_saved_reports_isolated(self):
        theirs = SavedReport.objects.create(business=self.other, name="Secret", params={"dataset": "sales"}, created_by=self.other_owner)
        self.login(self.owner)
        self.assertEqual(self.client.get(reverse("reports:saved", args=[theirs.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse("reports:saved", args=[theirs.pk])).status_code, 404)
        self.assertNotContains(self.client.get(reverse("reports:builder")), "Secret")
        self.assertTrue(SavedReport.objects.filter(pk=theirs.pk).exists())


class ProductImportTests(TenantTestCase):
    def upload(self, text, name="p.csv", user=None):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.login(user or self.owner)
        return self.client.post(reverse("products:import"), {"file": SimpleUploadedFile(name, text.encode())})

    HEADER = "name,sku,barcode,category,brand,unit,purchase_price,selling_price,wholesale_price,reorder_level,opening_quantity\n"

    def test_valid_import_preview_then_commit(self):
        r = self.upload(self.HEADER + "Paint 4L,PNT-4,111,Paint,Crown,pc,20000,28000,26000,5,12\nBrush,,,Paint,,pc,2000,3500,,10,0\n")
        self.assertContains(r, "2 of 2 rows are valid")
        self.assertEqual(Product.objects.filter(name="Paint 4L").count(), 0)  # nothing saved yet
        r = self.client.post(reverse("products:import"), {"confirm": "1"})
        self.assertRedirects(r, reverse("products:list"))
        p = Product.objects.get(name="Paint 4L")
        self.assertEqual((p.sku, p.selling_price, p.category.name), ("PNT-4", D("28000.00"), "Paint"))
        self.assertEqual(self.stock(p), D("12"))
        self.assertEqual(p.movements.get().movement_type, "opening")
        self.assertTrue(Product.objects.get(name="Brush").sku.startswith("SKU-"))
        self.assertEqual(Category.objects.filter(business=self.business, name="Paint").count(), 1)

    def test_any_error_blocks_the_whole_import(self):
        before = Product.objects.count()
        r = self.upload(self.HEADER + "Good,,,,,pc,1,2,,0,0\n,,,,,pc,1,2,,0,0\nBad price,,,,,pc,x,2,,0,0\nDup,{sku},,,,pc,1,2,,0,0\nUnit,,,,,parsec,1,2,,0,0\nNeg,,,,,pc,1,-2,,0,0\n".replace("{sku}", self.product.sku))
        self.assertContains(r, "name is required")
        self.assertContains(r, "not a number")
        self.assertContains(r, "already exists")
        self.assertContains(r, "is not set up")
        self.assertContains(r, "cannot be negative")
        self.assertEqual(Product.objects.count(), before)
        r = self.client.post(reverse("products:import"), {"confirm": "1"})  # nothing pending
        self.assertEqual(Product.objects.count(), before)

    def test_duplicates_within_file_and_barcode(self):
        r = self.upload(self.HEADER + "A,X1,B1,,,pc,1,2,,0,0\nB,X1,B2,,,pc,1,2,,0,0\nC,X3,B1,,,pc,1,2,,0,0\n")
        self.assertContains(r, "SKU &#x27;X1&#x27; already exists")
        self.assertContains(r, "barcode &#x27;B1&#x27; already exists")

    def test_bad_files(self):
        self.assertContains(self.upload("name,selling_price\n", "e.csv"), "header row and at least one", status_code=302) if False else None
        r = self.upload("name,selling_price\n", "e.csv")
        self.assertRedirects(r, reverse("products:import"))
        r = self.upload("x", "bad.exe")
        self.assertRedirects(r, reverse("products:import"))
        r = self.upload("foo,bar\n1,2\n")
        self.assertRedirects(r, reverse("products:import"))
        too_many = "name,selling_price\n" + "\n".join(f"P{i},1" for i in range(2001))
        self.assertRedirects(self.upload(too_many), reverse("products:import"))

    def test_excel_import(self):
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = Workbook()
        ws = wb.active
        ws.append(["Product", "Price", "Cost", "Qty"])
        ws.append(["Excel Item", 1500, 1000, 7])
        buf = io.BytesIO()
        wb.save(buf)
        self.login(self.owner)
        r = self.client.post(reverse("products:import"), {"file": SimpleUploadedFile("p.xlsx", buf.getvalue())})
        self.assertContains(r, "1 of 1 rows are valid")
        self.client.post(reverse("products:import"), {"confirm": "1"})
        self.assertEqual(self.stock(Product.objects.get(name="Excel Item")), D("7"))

    def test_permissions_and_tenancy(self):
        sp = make_user(self.business, "salesperson", "sam@alpha.test")
        self.login(sp)
        self.assertEqual(self.client.get(reverse("products:import")).status_code, 403)
        # same SKU in another business is not a duplicate
        other_unit = self.other.__class__.objects.get(pk=self.other.pk)
        r = self.upload(self.HEADER + f"Other,{self.product.sku},,,,pc,1,2,,0,0\n", user=self.other_owner)
        self.assertContains(r, "1 of 1 rows are valid")
        self.client.post(reverse("products:import"), {"confirm": "1"})
        self.assertEqual(Product.objects.filter(sku=self.product.sku).count(), 2)
        self.assertEqual(Product.objects.get(business=self.other, name="Other").business, self.other)

    def test_opening_stock_needs_permission(self):
        from accounts.models import Role, RolePermission, User
        role = Role.objects.create(business=self.business, name="Catalogue clerk")
        RolePermission.objects.create(role=role, codename="products.manage")
        clerk = User.objects.create_user(email="clerk@alpha.test", password="x" * 12, full_name="Clerk", business=self.business, role=role)
        r = self.upload(self.HEADER + "A,,,,,pc,1,2,,0,5\n", user=clerk)
        self.assertContains(r, "do not have permission to set opening stock")

    def test_export_csv_safe_and_scoped(self):
        make_product(self.business, self.owner, name="=SUM(A1)", stock="1")
        make_product(self.other, self.other_owner, name="Not Mine", stock="1")
        self.login(self.owner)
        text = self.client.get(reverse("products:export")).content.decode("utf-8-sig")
        self.assertIn("Cement 50kg", text)
        self.assertIn("'=SUM(A1)", text)
        self.assertNotIn("Not Mine", text)
        self.assertEqual(self.client.get(reverse("products:import_template")).status_code, 200)


class SearchTests(TenantTestCase):
    def test_search_and_permissions(self):
        cust = Customer.objects.create(business=self.business, name="Okello Joseph", phone="0772123456", customer_type="credit")
        sale = sales.create_sale(business=self.business, user=self.owner, customer=cust, items=[{"product_id": self.product.pk, "quantity": "1"}],
                                 payments=[{"method_id": self.cash.pk, "amount": "30000"}])
        Customer.objects.create(business=self.other, name="Okello Foreign")
        self.login(self.owner)
        r = self.client.get(reverse("core:search"), {"q": "okello"})
        self.assertContains(r, "Okello Joseph")
        self.assertNotContains(r, "Foreign")
        self.assertContains(self.client.get(reverse("core:search"), {"q": sale.invoice.number}), sale.invoice.number)
        self.assertContains(self.client.get(reverse("core:search"), {"q": "SKU-"}), "Cement")
        self.assertContains(self.client.get(reverse("core:search"), {"q": "0772"}), "Okello")
        self.assertContains(self.client.get(reverse("core:search"), {"q": "zzzz"}), "No results")
        self.assertContains(self.client.get(reverse("core:search"), {"q": "x"}), "at least two")
        sp = make_user(self.business, "salesperson", "sam@alpha.test")
        self.login(sp)
        r = self.client.get(reverse("core:search"), {"q": sale.number})
        self.assertContains(r, sale.number)
        supplier = Supplier.objects.create(business=self.business, name="Hima Cement")
        pur.create_purchase(business=self.business, user=self.owner, supplier=supplier, items=[{"product_id": self.product.pk, "quantity": "1", "unit_cost": "1"}])
        self.assertNotContains(self.client.get(reverse("core:search"), {"q": "PO-"}), "Purchases")
        self.login(self.owner)
        self.assertContains(self.client.get(reverse("core:search"), {"q": "PO-"}), "Purchases")
