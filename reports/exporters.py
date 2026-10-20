"""CSV / Excel / PDF renderers for ReportResult."""
import csv
import io
from datetime import date, datetime
from decimal import Decimal

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .engine import DATE, INT, MONEY_KIND, QTY

NAVY = colors.HexColor("#0b2a4a")
FORMULA_CHARS = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value):
    """Neutralise spreadsheet formula injection for user-controlled text."""
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(FORMULA_CHARS) else text


def display(value, kind, business):
    if value is None or value == "":
        return ""
    if kind == MONEY_KIND:
        return f"{Decimal(value):,.{business.currency_decimals}f}"
    if kind == QTY:
        text = f"{Decimal(value):,.3f}".rstrip("0").rstrip(".")
        return text or "0"
    if kind == "pct":
        return f"{Decimal(value):.1f}%"
    if kind == DATE and isinstance(value, (date, datetime)):
        return f"{value:%d %b %Y}"
    return str(value)


def to_csv(result, business):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([result.title])
    if result.subtitle:
        w.writerow([result.subtitle])
    w.writerow([c.label for c in result.columns])
    for row in result.rows:
        w.writerow([_raw(row.get(c.key), c.kind) for c in result.columns])
    if result.totals:
        w.writerow([("Total" if i == 0 else _raw(result.totals.get(c.key), c.kind)) for i, c in enumerate(result.columns)])
    return buf.getvalue().encode("utf-8-sig")


def _raw(value, kind):
    if value is None:
        return ""
    if kind in (MONEY_KIND, QTY, INT, "pct"):
        return str(value)
    if isinstance(value, (date, datetime)):
        return f"{value:%Y-%m-%d}"
    return safe_text(value)


def to_xlsx(result, business):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = result.title[:31].replace("/", "-")
    ws.append([result.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([result.subtitle])
    ws.append([])
    ws.append([c.label for c in result.columns])
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B2A4A")
    money_fmt = "#,##0" if business.currency_decimals == 0 else "#,##0." + "0" * business.currency_decimals

    def put(row_values):
        ws.append([_xl(v, c.kind) for v, c in zip(row_values, result.columns)])
        for cell, c in zip(ws[ws.max_row], result.columns):
            if c.kind == MONEY_KIND:
                cell.number_format = money_fmt
            elif c.kind == QTY:
                cell.number_format = "#,##0.###"
            elif c.kind == DATE:
                cell.number_format = "dd mmm yyyy"
            if c.kind in (MONEY_KIND, QTY, INT, "pct"):
                cell.alignment = Alignment(horizontal="right")

    for row in result.rows:
        put([row.get(c.key) for c in result.columns])
    if result.totals:
        put(["Total" if i == 0 else result.totals.get(c.key) for i, c in enumerate(result.columns)])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, c in enumerate(result.columns, start=1):
        width = max([len(c.label)] + [len(str(r.get(c.key, ""))) for r in result.rows[:200]]) + 2
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = min(max(width, 10), 45)
    ws.freeze_panes = "A5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xl(value, kind):
    if value is None or value == "":
        return None
    if kind in (MONEY_KIND, QTY, "pct"):
        return float(value)
    if kind == INT:
        return int(value)
    if kind == DATE:
        return value
    return safe_text(value)


def to_pdf(result, business):
    ss = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=ss["Title"], alignment=0, fontSize=16, textColor=NAVY, spaceAfter=2)
    small = ParagraphStyle("s", parent=ss["Normal"], fontSize=8, textColor=colors.HexColor("#64748b"))
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=7.5, leading=9)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm, title=result.title)
    story = [Paragraph(business.name, small), Paragraph(result.title, title), Paragraph(result.subtitle, small), Spacer(1, 4 * mm)]
    data = [[c.label for c in result.columns]]
    for row in result.rows:
        data.append([Paragraph(display(row.get(c.key), c.kind, business).replace("&", "&amp;").replace("<", "&lt;"), cell) if c.kind not in (MONEY_KIND, QTY, INT, "pct")
                     else display(row.get(c.key), c.kind, business) for c in result.columns])
    if result.totals:
        data.append(["Total" if i == 0 else display(result.totals.get(c.key), c.kind, business) for i, c in enumerate(result.columns)])
    table = Table(data, repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
             ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")])]
    for i, c in enumerate(result.columns):
        if c.kind in (MONEY_KIND, QTY, INT, "pct"):
            style.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
    if result.totals:
        style.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    story.append(table)
    if result.truncated:
        story.append(Paragraph("Only the first rows are shown. Narrow the filters to see everything.", small))
    if result.note:
        story.append(Spacer(1, 3 * mm))
        story.append(Paragraph(result.note, small))
    doc.build(story)
    return buf.getvalue()


EXPORTERS = {
    "csv": (to_csv, "text/csv; charset=utf-8", "csv"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": (to_pdf, "application/pdf", "pdf"),
}
