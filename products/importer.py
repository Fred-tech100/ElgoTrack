"""Product CSV/Excel import with full validation before anything is written."""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from core.utils import q2, q3
from inventory.models import StockMovement
from inventory.services import manual_adjustment

from .models import Category, Product, Unit

MAX_ROWS = 2000
MAX_BYTES = 2 * 1024 * 1024
COLUMNS = ["name", "sku", "barcode", "category", "brand", "unit", "purchase_price", "selling_price", "wholesale_price",
           "reorder_level", "opening_quantity", "description"]
ALIASES = {"product": "name", "product name": "name", "cost": "purchase_price", "cost price": "purchase_price", "buying price": "purchase_price",
           "price": "selling_price", "selling price": "selling_price", "wholesale": "wholesale_price", "wholesale price": "wholesale_price",
           "reorder": "reorder_level", "reorder level": "reorder_level", "opening stock": "opening_quantity", "opening quantity": "opening_quantity",
           "quantity": "opening_quantity", "qty": "opening_quantity", "unit of measure": "unit", "barcode": "barcode", "sku": "sku"}


class ImportFileError(Exception):
    pass


def read_rows(upload):
    if upload.size > MAX_BYTES:
        raise ImportFileError("The file is larger than 2 MB.")
    name = upload.name.lower()
    data = upload.read()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        except Exception:
            raise ImportFileError("This does not look like a valid Excel (.xlsx) file.")
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        raise ImportFileError("Upload a .csv or .xlsx file.")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise ImportFileError("The file needs a header row and at least one product.")
    if len(rows) - 1 > MAX_ROWS:
        raise ImportFileError(f"Too many rows (limit {MAX_ROWS}). Split the file and import in parts.")
    header = [ALIASES.get(h.strip().lower(), h.strip().lower().replace(" ", "_")) for h in rows[0]]
    if "name" not in header or "selling_price" not in header:
        raise ImportFileError("The header row must include at least 'name' and 'selling_price'.")
    return [(i + 2, dict(zip(header, [str(c).strip() for c in r]))) for i, r in enumerate(rows[1:])]


def _num(value, field, errors, default=None, places=2):
    if value == "":
        return default
    try:
        n = Decimal(value.replace(",", ""))
    except InvalidOperation:
        errors.append(f"{field} '{value}' is not a number")
        return default
    if n < 0:
        errors.append(f"{field} cannot be negative")
        return default
    return str(q2(n) if places == 2 else q3(n))


def validate(business, rows, can_set_opening):
    """Return (clean_rows, errors[(line, message)], new_categories). Nothing is saved."""
    units = {}
    for u in Unit.objects.filter(business=business):
        units[u.name.lower()] = u
        units[u.abbreviation.lower()] = u
    default_unit = units.get("piece") or next(iter(units.values()), None)
    cats = {c.name.lower() for c in Category.objects.filter(business=business)}
    db_skus = set(Product.objects.filter(business=business).values_list("sku", flat=True))
    db_barcodes = set(Product.objects.filter(business=business).exclude(barcode="").values_list("barcode", flat=True))
    seen_sku, seen_bar, clean, errors, new_cats = set(), set(), [], [], set()
    for line, r in rows:
        errs = []
        name = r.get("name", "")
        if not name:
            errs.append("name is required")
        if len(name) > 200:
            errs.append("name is longer than 200 characters")
        sku, barcode = r.get("sku", ""), r.get("barcode", "")
        if sku:
            if sku in db_skus or sku.lower() in {s.lower() for s in seen_sku}:
                errs.append(f"SKU '{sku}' already exists")
            seen_sku.add(sku)
        if barcode:
            if barcode in db_barcodes or barcode in seen_bar:
                errs.append(f"barcode '{barcode}' already exists")
            seen_bar.add(barcode)
        unit = default_unit
        if r.get("unit"):
            unit = units.get(r["unit"].lower())
            if unit is None:
                errs.append(f"unit '{r['unit']}' is not set up (add it under Categories & units)")
        cost = _num(r.get("purchase_price", ""), "purchase_price", errs, "0")
        price = _num(r.get("selling_price", ""), "selling_price", errs)
        if price is None and not errs:
            errs.append("selling_price is required")
        wholesale = _num(r.get("wholesale_price", ""), "wholesale_price", errs)
        reorder = _num(r.get("reorder_level", ""), "reorder_level", errs, "0", places=3)
        opening = _num(r.get("opening_quantity", ""), "opening_quantity", errs, "0", places=3)
        if opening not in (None, "0.000", "0") and not can_set_opening:
            errs.append("you do not have permission to set opening stock")
        cat = r.get("category", "")
        if cat and cat.lower() not in cats:
            new_cats.add(cat)
        if errs:
            errors.extend((line, e) for e in errs)
            continue
        clean.append({"name": name, "sku": sku, "barcode": barcode, "category": cat, "brand": r.get("brand", "")[:80],
                      "unit": unit.pk if unit else None, "purchase_price": cost, "selling_price": price, "wholesale_price": wholesale,
                      "reorder_level": reorder, "opening_quantity": opening, "description": r.get("description", "")})
    return clean, errors, sorted(new_cats)


@transaction.atomic
def commit(business, user, rows):
    categories = {c.name.lower(): c for c in Category.objects.filter(business=business)}
    created = 0
    for r in rows:
        category = None
        if r["category"]:
            category = categories.get(r["category"].lower())
            if category is None:
                category = categories[r["category"].lower()] = Category.objects.create(business=business, name=r["category"][:80])
        product = Product.objects.create(
            business=business, name=r["name"], sku=r["sku"], barcode=r["barcode"], category=category, brand=r["brand"],
            unit=Unit.objects.get(pk=r["unit"], business=business), purchase_price=Decimal(r["purchase_price"]),
            selling_price=Decimal(r["selling_price"]), wholesale_price=Decimal(r["wholesale_price"]) if r["wholesale_price"] else None,
            reorder_level=Decimal(r["reorder_level"]), description=r["description"])
        opening = Decimal(r["opening_quantity"])
        if opening > 0:
            manual_adjustment(business=business, user=user, product=product, movement_type=StockMovement.OPENING, quantity=opening, reason="Imported opening stock")
        created += 1
    return created
